from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\lin\Documents\ChatGPT\公路广告牌项目")
OUT = ROOT / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F3F5F7"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Microsoft YaHei", size=10.5, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def setup_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, before, after in (
        ("Title", 24, 0, 14),
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 13, 14, 6),
        ("Heading 3", 11, 10, 4),
    ):
        st = styles[style_name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = BLACK
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sp.add_run(subtitle)
    set_font(r, size=12)
    sp.paragraph_format.space_after = Pt(22)

    meta = [
        ("项目名称", "高速公路广告牌经营管理系统"),
        ("项目负责人", "林家明"),
        ("项目团队", "林家明（产品、UI）、李嘉伟（后端）、刘俊义（后端）、祝贺（测试）"),
        ("文档版本", "V1.0"),
        ("文档状态", "待客户评审"),
        ("基准日期", "2026 年 9 月 16 日"),
        ("提交对象", "客户"),
    ]
    add_table(doc, ["项目", "内容"], meta, widths=[3.2, 12.4])
    doc.add_page_break()
    add_heading(doc, "文档说明", 1)
    add_para(doc, "本文档以 2026 年 9 月 15 日完成的需求调研为基础，记录已经确认的产品、业务、设计或技术决策。客户单位正式名称、客户侧业务负责人和阿里云地域在发布前补齐。")
    return doc


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_font(r)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        r = p.add_run(f"{index}. {item}")
        set_font(r)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.text = str(text)
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Cm(widths[i])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_font(run, size=9.5, bold=True, color=RGBColor(255, 255, 255))
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            set_cell_borders(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Cm(widths[i])
            if idx % 2 == 1:
                set_cell_shading(cells[i], LIGHT_GRAY)
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 20 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_font(run, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_footer(doc, label):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"高速公路广告牌经营管理系统  {label}  V1.0")
        set_font(r, size=8, color=RGBColor(100, 100, 100))


def save(doc, filename, label):
    add_footer(doc, label)
    path = OUT / filename
    doc.save(path)
    return path


def build_prd():
    doc = setup_doc("高速公路广告牌经营管理系统产品需求文档", "明确首期范围 业务规则 验收标准与产品边界")
    add_heading(doc, "一 产品概述", 1)
    add_para(doc, "本产品面向高速公路沿线广告资源经营单位，统一管理广告牌资源、客户、租赁合同、应收实收、应付支出、发票、售后和经营分析。首期首先保证合同与资金业务闭环正确、账款可核对、状态可追溯，再逐步扩展智能问答和精细化分析。")
    add_heading(doc, "1 产品目标", 2)
    add_bullets(doc, [
        "所有有效广告资源、客户和合同均可在统一系统中检索与关联。",
        "应收、收款、核销、挂账、退款与应付分别记录并能够相互核对。",
        "合同、资金和合规事项的关键变更具备完整审计证据。",
        "核心经营报表直接由业务事实生成，不允许手工填写汇总金额。",
        "AI辅助合同录入和风险说明，但不直接改变正式业务记录。",
    ])
    add_heading(doc, "2 成功指标", 2)
    add_table(doc, ["指标", "目标"], [
        ("资源与合同覆盖", "上线范围内有效资源和有效合同 100% 可查询"),
        ("账款核对", "上线验收样本中应收 实收 核销与余额核对一致"),
        ("核心流程", "建档 签约 生成应收 收款核销 开票 报表完整通过验收"),
        ("合同智能录入", "固定样本集中关键字段准确率不低于 95%"),
        ("可追溯性", "关键合同与资金操作 100% 产生审计日志"),
    ], widths=[5, 10.6])

    add_heading(doc, "二 用户与权限", 1)
    add_table(doc, ["角色", "主要职责", "默认数据范围"], [
        ("运营档案员", "路段 点位 广告面 合规手续和照片维护", "指定运营单位"),
        ("合同管理员", "客户 合同 收费计划 变更 续签和解除", "指定运营单位"),
        ("财务管理员", "应收 收付款 核销 挂账 应付 发票和财务报表", "指定运营单位"),
        ("售后维护人员", "售后登记 处理 维修记录与费用", "本人或指定运营单位"),
        ("部门领导", "跨单位查看经营数据和风险", "授权单位或全部数据"),
        ("系统管理员", "账号 角色 配置和运行管理", "不自动获得业务数据权限"),
    ], widths=[3.2, 8.2, 4.2])
    add_para(doc, "一个用户可以拥有多个角色。功能权限取角色并集，数据访问仍受本人、指定运营单位或全部数据范围限制。导出、批量下载和敏感字段查看使用独立权限。")

    add_heading(doc, "三 产品范围", 1)
    add_table(doc, ["优先级", "范围"], [
        ("Must", "资源档案 客户 合同 收费计划 应收 收款核销 挂账 发票 核心报表 权限 审计 合同智能录入 风险摘要"),
        ("Should", "应付付款 售后 合同变更 提醒 经营收益测算 PDF导出"),
        ("Could", "自然语言经营问答 地图 移动端增强 滞纳金 扩展看板"),
        ("Won't", "客户门户 在线支付 外部财务系统对接 自动电子发票 原生 App 广告内容审核 完整会计核算"),
    ], widths=[3, 12.6])

    add_heading(doc, "四 核心业务需求", 1)
    modules = [
        ("1 广告资源管理", [
            "分别管理广告牌点位和可出租广告面。点位承载位置、产权与合规信息，广告面是合同租赁的最小资源。",
            "经营状态包括空闲、已预占、出租中、暂停经营和已拆除，由合同和停用信息共同计算。",
            "记录合同桩号、现场桩号和路政批复桩号，不因名称相近而覆盖。",
            "合规手续按类型记录编号、生效日、到期日、附件和状态。",
        ]),
        ("2 客户管理", [
            "维护客户编号、名称、联系人、联系方式、地址、开票和银行信息。",
            "黑名单客户不能创建新的生效合同，既有合同继续履行并进入风险提示。",
            "联系方式和银行信息在普通列表中遮挡，完整查看需要独立权限。",
        ]),
        ("3 合同管理", [
            "一份合同可以关联多个广告面，每个广告面可以具有独立租期。",
            "收费计划记录收费期间、应收金额和付款到期日，替代只能按自然年计费的年度价格。",
            "合同经过制单和不同人员确认后占用租赁时段并生成应收。",
            "已确认合同不得覆盖修改，通过变更单、续签或解除保留历史。",
        ]),
        ("4 应收与收款", [
            "合同确认后一次生成全部收费计划对应的应收，区分未到期、待收款、部分收款、已结清和逾期。",
            "收款流水与应收账单独立保存，支持一笔收款核销多条应收及一条应收由多笔收款核销。",
            "已确认客户但用途不明的款项进入客户挂账，付款人也不明确的款项进入待认领收款。",
            "多付款可继续挂账或通过退款单退回，退款关联原收款或挂账来源。",
        ]),
        ("5 应付 发票与售后", [
            "路政费、保险费、维护费等形成独立应付和付款记录，不与应收台账混合。",
            "发票记录可以关联客户、合同、应收和收款，首期不对接税务开票系统。",
            "售后事项关联客户、合同或广告资源，保存责任人、状态、处理记录、费用和结果。",
        ]),
        ("6 AI 能力", [
            "合同 PDF 或图片经过文本提取或 OCR 后生成结构化合同草稿，并显示字段来源页码或原文片段。",
            "到期、逾期和合规事实由确定性规则计算，AI负责整理风险摘要与处理建议。",
            "文档问答只能检索用户有权访问的合同、制度、政策和操作手册，回答必须标注来源。",
            "模型通过可配置 API Key 接入兼容 OpenAI 协议的国产大模型，测试时按效果选用通义或 DeepSeek。",
        ]),
    ]
    for title, bullets in modules:
        add_heading(doc, title, 2)
        add_bullets(doc, bullets)

    add_heading(doc, "五 关键状态与规则", 1)
    add_table(doc, ["对象", "状态或规则"], [
        ("广告面", "空闲 已预占 出租中 暂停经营 已拆除"),
        ("合同", "草稿 待确认 已生效 已到期 已解除 已作废"),
        ("应收", "未到期 待收款 部分收款 已结清 逾期 已调整"),
        ("售后", "待处理 处理中 待验收 已完成 已关闭"),
        ("删除", "正式合同和资金记录不得删除，错误通过作废 冲销或调整处理"),
        ("金额", "定点数值保留两位小数，核销不得超过应收未收金额"),
    ], widths=[3.5, 12.1])

    add_heading(doc, "六 报表与指标", 1)
    add_table(doc, ["报表", "主要指标", "口径"], [
        ("资源经营", "空闲 预占 出租中 暂停经营 拆除数量", "广告面状态"),
        ("合同执行", "待开始 执行中 即将到期 已到期 已解除", "合同与租期"),
        ("应收欠款", "应收 已核销 未收 逾期与账龄", "收费期间和到期日"),
        ("资金收支", "收款 付款 退款 现金净流入", "实际资金日期"),
        ("经营收益测算", "租金应收减已确认经营成本", "不等同会计利润"),
    ], widths=[3.2, 7.0, 5.4])

    add_heading(doc, "七 非功能需求", 1)
    add_table(doc, ["类别", "要求"], [
        ("性能", "普通查询 95% 在 2 秒内返回，普通写操作 95% 在 3 秒内完成，常用报表 10 秒内生成"),
        ("可用性", "正式环境月度可用性目标 99.5%，AI Redis或Milvus故障不得阻断核心合同财务业务"),
        ("安全", "HTTPS 强密码哈希 服务端权限校验 参数化查询 私有存储 敏感字段遮挡和最小权限"),
        ("审计", "关键操作记录用户 角色 时间 对象 变更内容 原因 IP 结果及确认信息"),
        ("兼容", "支持最新版及前一主要版本 Chrome和Edge，1366×768可完成全部核心操作"),
        ("备份", "RDS每日备份，RPO不超过24小时，RTO不超过4小时，每季度恢复演练"),
    ], widths=[3, 12.6])

    add_heading(doc, "八 上线与验收", 1)
    add_bullets(doc, [
        "PRD、产品设计方案和原型经客户确认。",
        "Must范围完成并通过核心业务、权限矩阵和自动化测试。",
        "迁移数据经客户核对，备份恢复演练成功。",
        "无阻断级或严重级缺陷，操作手册与管理员手册完成。",
        "客户正式数据进入系统前，FastAPI迁移至阿里云ECS并使用内网资源。",
    ])
    add_heading(doc, "九 待补充信息", 1)
    add_bullets(doc, ["客户单位正式名称", "客户侧业务负责人", "阿里云正式环境地域", "客户确认的合同样本和历史数据模板"])
    return save(doc, "01_高速公路广告牌经营管理系统_PRD_V1.0.docx", "产品需求文档")


def build_product_design():
    doc = setup_doc("高速公路广告牌经营管理系统产品设计方案", "定义信息架构 页面结构 核心交互与状态呈现")
    add_heading(doc, "一 设计目标", 1)
    add_para(doc, "产品设计以高频后台操作的准确性和可核对性为优先，重点降低合同录入、账款核销和风险追踪的认知负担。界面使用一致的列表、详情、抽屉和确认流程，AI结果与正式业务数据保持清晰区别。")
    add_heading(doc, "二 信息架构", 1)
    add_table(doc, ["一级导航", "二级页面"], [
        ("工作台", "待办提醒 风险摘要 核心指标 最近操作"),
        ("资源管理", "路段 点位 广告面 合规手续 地图后续扩展"),
        ("客户管理", "客户列表 客户详情 黑名单 合作记录"),
        ("合同管理", "合同列表 新建合同 待确认 变更续签解除 收费计划"),
        ("财务管理", "应收 收款核销 挂账 待认领 退款 应付付款 发票"),
        ("售后管理", "售后事项 处理任务 维修记录"),
        ("经营分析", "资源经营 合同执行 应收欠款 资金收支 经营收益"),
        ("AI助手", "合同智能录入 文档问答 风险摘要 AI任务记录"),
        ("系统管理", "用户 角色 权限 运营单位 字典 审计日志"),
    ], widths=[3.5, 12.1])

    add_heading(doc, "三 页面设计规范", 1)
    add_bullets(doc, [
        "列表页固定包含筛选区、操作区、数据表格、分页和列配置；复杂筛选可展开，不长期占用首屏。",
        "详情页使用摘要、关联信息和时间线组合，正式记录不提供直接覆盖修改入口。",
        "新增和低复杂度编辑使用抽屉，合同、核销和变更等复杂任务使用独立页面。",
        "危险操作使用二次确认，并要求填写原因；退款、调整和作废进入双人确认。",
        "金额、状态、到期日和风险等级在所有页面使用一致的格式和颜色语义。",
        "AI生成字段显示待确认标识、来源片段和置信提示，不与人工确认字段混淆。",
    ])

    add_heading(doc, "四 核心页面清单", 1)
    add_table(doc, ["页面", "主要区域", "关键操作"], [
        ("工作台", "待办 风险 指标 趋势 最近操作", "进入待确认合同 逾期账单或合规事项"),
        ("广告面列表", "路段 状态 租赁期 合规风险", "建档 停用 查看租赁历史"),
        ("合同编辑", "客户 广告面 租期 收费计划 附件", "保存草稿 提交确认 AI导入"),
        ("合同详情", "摘要 资源 收费 应收 变更 时间线", "变更 续签 解除 下载附件"),
        ("收款核销", "收款信息 未结应收 分配结果", "自动建议 手工调整 确认核销"),
        ("客户资金详情", "应收 收款 核销 挂账 退款 发票", "下钻来源 导出对账"),
        ("合同智能录入", "文件预览 提取字段 来源片段", "重新识别 修改草稿 确认创建"),
        ("经营报表", "筛选 指标 图表 明细", "下钻 导出 保存筛选"),
    ], widths=[3.5, 7.3, 4.8])

    add_heading(doc, "五 核心交互流程", 1)
    flows = [
        ("合同创建", "选择或创建客户 → 选择广告面和租期 → 检查时间冲突 → 配置收费计划 → 上传附件 → 保存草稿 → 提交另一人确认 → 占用租期并生成应收"),
        ("智能录入", "上传合同 → 文本提取或OCR → 模型结构化输出 → 规则校验 → 展示来源与待确认字段 → 用户修正 → 创建合同草稿"),
        ("收款核销", "登记收款 → 识别客户 → 展示未结应收 → 默认优先最早到期 → 财务调整分配 → 校验金额 → 确认核销"),
        ("未知收款", "登记流水 → 无法确认客户 → 进入待认领 → 后续确认客户 → 转客户挂账或直接核销"),
        ("合同变更", "发起变更 → 比较前后内容 → 计算应收影响 → 提交确认 → 生成调整记录 → 更新未来租赁状态"),
    ]
    add_table(doc, ["流程", "步骤"], flows, widths=[3.2, 12.4])

    add_heading(doc, "六 状态呈现", 1)
    add_table(doc, ["对象", "主状态", "界面要求"], [
        ("广告面", "空闲 预占 出租中 暂停经营 拆除", "状态由系统计算时禁止手工覆盖，并显示形成原因"),
        ("合同", "草稿 待确认 已生效 已到期 已解除 已作废", "详情页显示状态时间线及责任人"),
        ("应收", "未到期 待收 部分收款 已结清 逾期", "金额与到期日并列，逾期显示账龄"),
        ("AI任务", "排队 处理中 待确认 已完成 失败", "允许离开页面后查看，失败提供可理解原因"),
    ], widths=[3.2, 6.0, 6.4])

    add_heading(doc, "七 权限与敏感操作体验", 1)
    add_bullets(doc, [
        "无权操作的按钮不显示；通过直接请求访问时服务端仍返回明确的无权限结果。",
        "完整银行账户、批量导出和附件批量下载使用独立权限。",
        "双人确认页面必须显示制单人、差异、附件和操作原因。",
        "审计日志提供对象、操作人、时间、变更字段和请求追踪编号筛选。",
    ])

    add_heading(doc, "八 原型交付要求", 1)
    add_bullets(doc, [
        "制作可点击高保真原型，至少覆盖登录、工作台、广告面建档、合同创建确认、收款核销、合同智能录入和经营报表。",
        "原型以 1440×900 为基准并验证 1366×768，使用统一组件和状态样例。",
        "每个核心流程同时包含正常、空数据、校验失败、无权限和依赖故障状态。",
        "原型版本与 PRD 版本绑定，客户确认后再进入对应 Sprint。",
    ])
    return save(doc, "02_高速公路广告牌经营管理系统_产品设计方案_V1.0.docx", "产品设计方案")


def build_domain():
    doc = setup_doc("高速公路广告牌经营管理系统业务流程与领域模型设计", "统一业务语言 对象关系 状态和不变量")
    add_heading(doc, "一 领域边界", 1)
    add_para(doc, "系统围绕资源、客户、合同、资金、售后、经营分析和身份权限七个业务模块组织。模块通过稳定接口协作，正式业务事实集中保存在 MySQL；Redis和Milvus只提供加速或检索能力。")
    add_heading(doc, "二 核心对象关系", 1)
    add_table(doc, ["对象", "关系", "约束"], [
        ("路段与点位", "一个路段包含多个广告牌点位", "点位必须属于一个运营单位"),
        ("点位与广告面", "一个点位包含一个或多个广告面", "广告面是最小可出租资源"),
        ("合同与广告面", "多对多，通过合同广告面明细关联", "同一广告面生效租期不得重叠"),
        ("合同与收费计划", "一个合同包含多期收费计划", "每期具有起止日期 金额和到期日"),
        ("收费计划与应收", "合同确认后形成正式应收", "应收历史不得覆盖"),
        ("收款与应收", "通过核销明细形成多对多关系", "核销总额不得超过两侧可用余额"),
        ("客户与挂账", "已确认客户的未分配收款余额", "客户未知时必须进入待认领"),
        ("合同与发票", "发票可关联合同 应收和收款", "首期只登记不自动开票"),
    ], widths=[3.4, 6.2, 6.0])

    add_heading(doc, "三 核心不变量", 1)
    add_numbered(doc, [
        "同一广告面不能存在时间重叠的两份已确认合同。",
        "合同制单人与确认人不得为同一用户。",
        "合同确认后才占用租赁时段并生成应收。",
        "应收、收款、核销、挂账、退款和调整分别保存。",
        "单次核销不得超过应收未收金额或收款未分配金额。",
        "已核销或已开票的业务记录不得物理删除。",
        "保证金等暂收款不得计入租金收入。",
        "AI建议必须经过业务规则校验和用户确认才能形成正式记录。",
        "用户通过AI访问的数据不得超出其原有功能权限和数据范围。",
    ])

    add_heading(doc, "四 状态生命周期", 1)
    add_table(doc, ["对象", "状态链", "异常出口"], [
        ("合同", "草稿 → 待确认 → 已生效 → 已到期", "待确认可驳回；已生效可解除；错误记录可作废"),
        ("应收", "未到期 → 待收款 → 部分收款 → 已结清", "到期未结清进入逾期；金额变化产生调整"),
        ("售后事项", "待处理 → 处理中 → 待验收 → 已完成", "无须处理或重复事项可以关闭"),
        ("AI任务", "排队 → 处理中 → 待确认 → 已完成", "失败可重试，不能留下半完成正式数据"),
    ], widths=[3.2, 7.2, 5.2])

    add_heading(doc, "五 业务流程", 1)
    add_heading(doc, "1 资源到合同", 2)
    add_numbered(doc, ["创建路段、点位和广告面档案。", "补充产权、用地和合规手续。", "客户建档并检查黑名单。", "合同选择广告面和租期，系统检查冲突。", "配置收费计划并提交确认。", "另一名用户确认后占用租期并生成应收。"])
    add_heading(doc, "2 收款到核销", 2)
    add_numbered(doc, ["登记真实收款流水。", "确认付款客户及款项类型。", "展示客户未结应收并默认优先最早到期。", "财务调整核销分配并提交。", "事务内同时更新核销明细和余额。", "无法识别客户时进入待认领，客户明确但用途不明时进入挂账。"])
    add_heading(doc, "3 合同变更", 2)
    add_numbered(doc, ["从已生效合同发起变更单。", "记录租期、资源、收费或附件的前后差异。", "计算应收、开票和资源占用影响。", "另一名有权限用户确认。", "生成调整和新历史版本，不覆盖原记录。"])
    add_heading(doc, "4 AI合同录入", 2)
    add_numbered(doc, ["附件上传至私有OSS。", "根据文件类型执行文本提取或OCR。", "模型按Pydantic定义输出结构化字段。", "系统执行日期、金额、资源和必填规则校验。", "用户对照来源片段修正并确认。", "仅创建合同草稿，不直接生效。"])

    add_heading(doc, "六 术语表", 1)
    terms = [
        ("广告牌点位", "具有固定位置、产权及合规资料的实体设施。"),
        ("广告面", "点位上能够独立签约出租的最小经营资源。"),
        ("收费计划", "合同约定的分期收费安排。"),
        ("应收账单", "客户在指定到期日前应承担的正式付款义务。"),
        ("收款流水", "运营方实际收到款项的独立资金记录。"),
        ("核销", "把收款金额分配至具体应收的过程。"),
        ("客户挂账", "客户已明确但具体用途尚未明确的收款余额。"),
        ("待认领收款", "付款客户身份仍不明确的收款。"),
        ("经营收益测算", "期间租金应收减已确认经营成本，不等同会计利润。"),
        ("AI建议", "等待用户确认的模型输出，不是正式业务记录。"),
    ]
    add_table(doc, ["术语", "定义"], terms, widths=[4, 11.6])
    return save(doc, "03_高速公路广告牌经营管理系统_业务流程与领域模型设计_V1.0.docx", "业务流程与领域模型设计")


def build_architecture():
    doc = setup_doc("高速公路广告牌经营管理系统技术架构设计", "定义技术选型 模块接口 数据安全 部署与演进方案")
    add_heading(doc, "一 架构结论", 1)
    add_para(doc, "系统采用 Vue 3 SPA 与 FastAPI 模块化单体的前后端分离架构。个人开发和试用阶段可部署在 Vercel；客户正式使用前迁移至阿里云，前端使用 OSS 与 CDN，FastAPI 和后台任务以 Docker 运行在 ECS，RDS MySQL、Redis、OSS和Milvus使用托管服务。")
    add_heading(doc, "二 技术栈", 1)
    add_table(doc, ["层次", "技术选型"], [
        ("前端", "Vue 3 Vite TypeScript Vue Router Pinia Element Plus Axios TanStack Query Zod"),
        ("后端", "FastAPI Pydantic v2 SQLAlchemy 2 Alembic asyncmy uv"),
        ("业务数据", "阿里云 RDS MySQL 8.0"),
        ("缓存", "阿里云 Redis"),
        ("文件", "阿里云 OSS 私有 Bucket"),
        ("AI", "兼容 OpenAI 协议的国产大模型 API Key LangChain LangGraph按需"),
        ("向量检索", "本地 Milvus Lite或Docker Milvus 生产Milvus托管版"),
        ("测试", "pytest Vitest Playwright"),
        ("部署", "GitHub Actions Docker Vercel试用 阿里云正式环境"),
        ("可观测", "结构化日志 Prometheus Grafana Langfuse"),
    ], widths=[3.7, 11.9])

    add_heading(doc, "三 模块设计", 1)
    add_table(doc, ["模块", "职责", "主要接口"], [
        ("资源", "路段 点位 广告面 合规手续和经营状态", "创建档案 查询可用资源 计算经营状态"),
        ("客户", "客户资料 黑名单 开票信息和敏感字段", "维护客户 校验签约资格"),
        ("合同", "合同 收费计划 确认 变更 续签 解除", "创建草稿 确认合同 执行变更"),
        ("应收与核销", "应收 收款 核销 挂账 待认领 退款 调整", "生成应收 登记收款 分配核销 发起退款"),
        ("应付与发票", "支出 付款和发票台账", "登记应付付款 发票关联"),
        ("经营分析", "指标计算 报表下钻和导出", "查询指标 获取明细 创建导出"),
        ("AI应用", "合同提取 风险摘要 文档检索和成本记录", "解析合同 生成摘要 检索文档"),
        ("身份权限", "用户 角色 数据范围 会话和审计", "鉴权 授权 记录审计"),
    ], widths=[3.2, 7.0, 5.4])

    add_heading(doc, "四 外部依赖接口与适配器", 1)
    add_table(doc, ["接口", "生产适配器", "测试适配器", "不变量"], [
        ("对象存储", "AlibabaOSSAdapter", "LocalStorageAdapter", "私有访问和业务权限先于下载"),
        ("向量存储", "MilvusAdapter", "LocalMilvusAdapter", "删除文档时索引同步失效"),
        ("模型网关", "OpenAICompatibleModelAdapter", "FakeModelAdapter", "模型输出只能形成建议"),
        ("缓存", "RedisAdapter", "InMemoryCacheAdapter", "缓存不可成为唯一事实来源"),
        ("OCR", "可配置OCRAdapter", "FixtureOCRAdapter", "保留页码和来源片段"),
    ], widths=[3.0, 4.3, 4.0, 4.3])

    add_heading(doc, "五 数据设计原则", 1)
    add_bullets(doc, [
        "正式合同、应收、收款、核销、退款和调整使用独立表保存，不通过覆盖更新伪造历史。",
        "金额使用 DECIMAL，不在 JavaScript 或 Python 中使用二进制浮点进行账款计算。",
        "核销、退款和合同确认使用数据库事务与必要的行锁，防止并发超额分配。",
        "所有业务表包含稳定主键、创建与更新时间、创建与更新用户以及乐观并发字段。",
        "数据库结构变更只能通过 Alembic迁移，生产迁移采用先兼容后清理策略。",
    ])

    add_heading(doc, "六 AI与RAG架构", 1)
    add_numbered(doc, [
        "附件写入OSS并在MySQL创建AI任务。",
        "后台任务执行文本提取或OCR，保存文档页和来源定位。",
        "合同提取直接使用结构化输出和Pydantic校验，不强制经过LangChain。",
        "需要文档问答时执行切分、嵌入并写入Milvus，同时保存文档与向量关联。",
        "检索前执行用户、运营单位和文档权限过滤，回答附带来源。",
        "通过Langfuse记录链路、模型、耗时、token和结果评价，不记录不必要的敏感正文。",
    ])

    add_heading(doc, "七 安全设计", 1)
    add_bullets(doc, [
        "密码使用Argon2id等可靠算法，访问令牌和会话使用安全Cookie或等价安全机制。",
        "权限在FastAPI服务端校验，前端路由和按钮控制只作为用户体验增强。",
        "RDS Redis和Milvus在正式环境仅通过VPC内网访问，OSS使用私有Bucket。",
        "API Key、数据库密码和云凭据保存在环境变量或密钥管理服务，不进入Git。",
        "AI工具调用继承当前用户权限，系统管理员也不自动获得全部业务数据权限。",
    ])

    add_heading(doc, "八 部署与CI CD", 1)
    add_table(doc, ["阶段", "前端", "后端", "数据与限制"], [
        ("个人试用", "Vercel Vue Vite", "Vercel FastAPI Python Function", "只使用测试或脱敏数据，保留Dockerfile"),
        ("正式生产", "阿里云OSS与CDN", "ECS Docker FastAPI与任务进程", "RDS Redis OSS Milvus同地域VPC"),
    ], widths=[3, 4.1, 4.5, 4.0])
    add_para(doc, "GitHub Actions执行静态检查、类型检查、单元测试、集成测试、前端构建和镜像构建。功能分支产生预览环境；主分支发布需经过测试门禁。生产数据库迁移作为受控发布步骤执行，不能在每次请求中自动运行。")

    add_heading(doc, "九 可观测与恢复", 1)
    add_bullets(doc, [
        "每个请求生成追踪编号，日志记录模块、业务对象、耗时和结果，但不记录密码、令牌和完整敏感正文。",
        "监控接口错误率、延迟、数据库连接、后台任务失败、AI耗时、token和模型错误。",
        "RDS每日备份并支持时间点恢复，OSS启用版本或防误删，Milvus可从原文和切分记录重建。",
        "AI、Redis或Milvus故障时执行降级，合同和财务核心流程继续运行。",
    ])
    add_heading(doc, "十 架构图", 1)
    add_para(doc, "随本文档一并交付可交互架构图《系统目标架构图.html》，展示业务请求、合同AI处理、正式事实来源和部署演进。")
    return save(doc, "04_高速公路广告牌经营管理系统_技术架构设计_V1.0.docx", "技术架构设计")


def build_test():
    doc = setup_doc("高速公路广告牌经营管理系统测试与验收方案", "定义测试层次 核心用例 质量门禁与上线准入")
    add_heading(doc, "一 测试目标", 1)
    add_para(doc, "测试重点是证明合同租赁时段、应收实收、核销余额、权限隔离和AI辅助边界正确，而不是只验证页面能够打开。验收以真实业务场景、固定数据集和可复核结果为依据。")
    add_heading(doc, "二 测试层次", 1)
    add_table(doc, ["层次", "工具", "重点"], [
        ("静态检查", "Ruff Pyright ESLint TypeScript", "格式 类型 安全规则和常见缺陷"),
        ("单元测试", "pytest Vitest", "金额 状态 时间重叠 权限和风险规则"),
        ("集成测试", "pytest MySQL Redis OSS Milvus测试环境", "事务 约束 适配器和故障降级"),
        ("接口测试", "pytest OpenAPI", "输入校验 错误码 权限 幂等与分页"),
        ("端到端测试", "Playwright", "登录 建档 签约 核销 AI录入 报表和导出"),
        ("验收测试", "客户业务样本", "业务口径 数据核对和可操作性"),
    ], widths=[3, 4.4, 8.2])

    add_heading(doc, "三 核心验收场景", 1)
    rows = [
        ("AT-01", "广告资源建档", "创建点位和两个广告面，补充三个桩号及合规手续", "资源可查，状态与合规风险正确"),
        ("AT-02", "多广告面合同", "同一合同关联多个广告面并设置不同收费计划", "确认后占用租期并生成对应应收"),
        ("AT-03", "租期冲突", "为已占用广告面创建重叠合同", "系统阻止确认并指出冲突合同"),
        ("AT-04", "部分收款", "一条应收由两笔收款完成", "未收金额逐次减少，最终结清"),
        ("AT-05", "合并收款", "一笔收款核销多条应收", "各应收和收款未分配余额正确"),
        ("AT-06", "未知收款", "付款客户无法确认", "进入待认领，不影响任何客户余额"),
        ("AT-07", "合同变更", "已生效合同调整未来收费", "原记录保留并生成变更及应收调整"),
        ("AT-08", "退款", "从客户挂账发起退款", "双人确认后生成退款并更新余额"),
        ("AT-09", "权限隔离", "A运营单位用户查询B单位合同", "页面和接口均拒绝且AI不泄露信息"),
        ("AT-10", "合同智能录入", "上传样本合同并确认提取结果", "字段带来源，规则校验通过后只生成草稿"),
        ("AT-11", "风险摘要", "准备到期 欠款和证件过期数据", "规则命中准确，AI摘要引用来源记录"),
        ("AT-12", "报表下钻", "查看逾期金额并下钻", "汇总与明细一致，可定位应收和合同"),
    ]
    add_table(doc, ["编号", "场景", "输入或操作", "预期结果"], rows, widths=[2, 3.2, 5.3, 5.1])

    add_heading(doc, "四 AI专项验收", 1)
    add_table(doc, ["项目", "验收方法", "标准"], [
        ("合同提取准确率", "使用客户确认的固定样本集逐字段比对", "关键字段准确率不低于95%"),
        ("不确定性处理", "准备模糊 缺失和冲突字段", "留空或标记待确认，不猜测"),
        ("来源可追溯", "抽查提取字段", "展示正确页码或原文片段"),
        ("权限", "跨单位提问和检索", "不返回无权文档 名称 数量或摘要"),
        ("有据率", "固定RAG问题集", "回答带来源，无法确认时明确拒答"),
        ("成本", "统计各功能调用", "记录用户 模型 token 耗时和费用标签"),
        ("降级", "模拟模型或Milvus不可用", "AI功能提示故障，核心业务仍可使用"),
    ], widths=[3.2, 6.2, 6.2])

    add_heading(doc, "五 性能与可靠性", 1)
    add_table(doc, ["项目", "目标"], [
        ("普通查询", "20人并发下95%请求2秒内返回"),
        ("普通写操作", "95%请求3秒内完成"),
        ("常用报表", "10秒内生成"),
        ("Excel校验", "5000行在60秒内给出结果或进入后台任务"),
        ("AI流式响应", "10秒内出现首个响应"),
        ("合同解析", "普通合同2分钟内完成，可后台查看"),
        ("恢复演练", "RPO不超过24小时，RTO不超过4小时"),
    ], widths=[5, 10.6])

    add_heading(doc, "六 安全测试", 1)
    add_bullets(doc, [
        "验证未登录、越权、跨运营单位、参数篡改和直接调用接口场景。",
        "验证登录失败锁定、速率限制、会话失效和管理员停用账号。",
        "验证SQL注入、文件伪装、超大文件、恶意文件名和不安全下载。",
        "验证日志、导出和AI调用中不出现密码、令牌和不必要的敏感正文。",
        "验证系统管理员在未授予业务角色时不能查看合同和财务数据。",
    ])

    add_heading(doc, "七 缺陷等级", 1)
    add_table(doc, ["等级", "定义", "上线要求"], [
        ("阻断级", "系统无法使用 数据损坏 严重资金错误或重大权限泄露", "必须清零"),
        ("严重级", "核心流程失败或重要业务结果错误且无可接受绕行", "必须清零"),
        ("一般级", "非核心功能错误或存在可接受绕行", "经评估后限期修复"),
        ("轻微级", "文案 样式或低影响体验问题", "记录并排期"),
    ], widths=[3, 8.2, 4.4])

    add_heading(doc, "八 上线准入", 1)
    add_numbered(doc, [
        "PRD 产品设计方案 原型和验收范围经客户确认。",
        "Must需求完成，所有核心验收场景通过。",
        "单元 集成 接口和端到端测试达到团队约定门禁。",
        "阻断级和严重级缺陷清零。",
        "权限矩阵 数据迁移和账款余额完成专项核对。",
        "备份恢复演练成功，操作手册和管理员手册完成。",
        "正式客户数据使用阿里云正式环境和内网托管资源。",
        "客户侧业务负责人签署上线确认。",
    ])
    return save(doc, "05_高速公路广告牌经营管理系统_测试与验收方案_V1.0.docx", "测试与验收方案")


if __name__ == "__main__":
    paths = [build_prd(), build_product_design(), build_domain(), build_architecture(), build_test()]
    for p in paths:
        print(p)
